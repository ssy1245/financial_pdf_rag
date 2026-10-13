"""Native PDF/DOCX/XLSX adapters; preserve the structured original before retrieval views."""
import hashlib
from pathlib import Path
from datetime import date, datetime, time
from zipfile import ZipFile
from .document import Document, Element, Table, Cell, SourceLocation, OCRRequest


def parse_document(path: Path, filename: str | None = None) -> Document:
    path = Path(path)
    fmt = path.suffix.lower().lstrip(".")
    if fmt not in {"pdf", "docx", "xlsx"}:
        raise ValueError("仅支持 .pdf、.docx、.xlsx；旧版 .doc/.xls 请先转换")
    if fmt != "pdf":
        with ZipFile(path) as archive:
            if len(archive.infolist()) > 20000 or sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("Office 文档解压体积过大")
    doc = Document(document_id=hashlib.sha256(path.read_bytes()).hexdigest(), filename=filename or path.name, format=fmt)
    {"pdf": _pdf, "docx": _word, "xlsx": _excel}[fmt](path, doc)
    return Document.model_validate(doc.model_dump())


def _add_table(doc, rows, columns, cells, source, **kwargs):
    tid = f"t{len(doc.tables)+1}"
    table = Table(table_id=tid, rows=rows, columns=columns, cells=cells, source=source, **kwargs)
    doc.tables.append(table)
    doc.elements.append(Element(element_id=f"e{len(doc.elements)+1}", kind="table", table_id=tid, source=source))


def _pdf(path, doc):
    import pymupdf
    with pymupdf.open(path) as pdf:
        if pdf.needs_pass:
            raise ValueError("加密 PDF 需要先解密")
        doc.page_count = len(pdf)
        for number, page in enumerate(pdf, 1):
            tables = []
            try:
                tables = list(page.find_tables().tables)
            except Exception:
                doc.warnings.append(f"第 {number} 页表格检测失败，保留原生文字待核验")
            table_boxes = [pymupdf.Rect(t.bbox) for t in tables]
            # Use line-level boxes: a text block may straddle a detected table.
            items = []
            for block in page.get_text("dict", sort=True)["blocks"]:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    box = pymupdf.Rect(line["bbox"])
                    if not text or any(b.contains((box.tl + box.br) / 2) for b in table_boxes):
                        continue
                    items.append((box.y0, box.x0, text, list(box)))
            for y, x, text, box in items:
                doc.elements.append(Element(element_id=f"e{len(doc.elements)+1}", kind="paragraph", text=text,
                    source=SourceLocation(format="pdf", page=number, bbox=box)))
            for index, table in enumerate(tables, 1):
                source = SourceLocation(format="pdf", page=number, bbox=list(table.bbox), table=index)
                texts = table.extract()
                # PyMuPDF's grid is defined by cell origins, not every right/bottom edge.
                # Uneven row widths otherwise introduce phantom columns.
                xs = sorted({b[0] for row in table.rows for b in row.cells if b})
                ys = sorted({b[1] for row in table.rows for b in row.cells if b})
                cells = []
                for r, row in enumerate(table.rows, 1):
                    for c, box in enumerate(row.cells, 1):
                        if box is None:
                            continue  # covered by merged anchor, not an empty cell
                        rs = max(1, sum(box[1] <= y < box[3]-1e-6 for y in ys))
                        cs = max(1, sum(box[0] <= x < box[2]-1e-6 for x in xs))
                        cells.append(Cell(row=r, column=c, row_span=rs, column_span=cs,
                            text=texts[r-1][c-1] or "", source=source.model_copy(update={"bbox":list(box),"row_start":r,"row_end":r+rs-1,"column_start":c,"column_end":c+cs-1})))
                _add_table(doc, table.row_count, table.col_count, cells, source,
                    header_rows=[] if table.header.external else [1], header_origin="heuristic",
                    column_names=[str(n or "") for n in table.header.names],
                    column_header_sources=[source.model_copy(update={"bbox":list(box)}) for box in table.header.cells if box])
            # Native text does not establish that all image regions were read.
            images = page.get_image_info()
            if not page.get_text().strip():
                doc.ocr_requests.append(OCRRequest(document_id=doc.document_id, page=number,
                    bbox=list(page.rect), reason="页面无原生文字，需核验空白页或扫描页"))
            else:
                for img in images:
                    box = pymupdf.Rect(img["bbox"])
                    if box.get_area() > page.rect.get_area() * 0.15:
                        doc.ocr_requests.append(OCRRequest(document_id=doc.document_id,page=number,
                            bbox=list(box),reason="大幅图像区域未识别，可能包含扫描文字或表格"))
        doc.elements.sort(key=lambda e:(e.source.page, e.source.bbox[1], e.source.bbox[0]))
    doc.warnings.append("PDF 表格为原生线框检测；无框表、复杂多栏及跨页续表需核验，跨页不自动合并")


def _word(path, doc):
    from docx import Document as WordDocument
    from docx.text.paragraph import Paragraph
    word = WordDocument(path)
    headings, paragraph_index, table_index = [], 0, 0
    for block in word.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            style = block.style.name if block.style else ""
            heading = style.startswith("Heading")
            if heading:
                try: level = int(style.split()[-1])
                except ValueError: level = 1
                headings = headings[:level-1] + [block.text]
            if block.text.strip():
                doc.elements.append(Element(element_id=f"e{len(doc.elements)+1}", kind="heading" if heading else "paragraph",
                    text=block.text, source=SourceLocation(format="docx",paragraph=paragraph_index,heading_path=list(headings))))
        else:
            table_index += 1
            source = SourceLocation(format="docx",table=table_index,heading_path=list(headings))
            cells, anchors, headers = [], {}, []
            for r, row in enumerate(block.rows, 1):
                if row._tr.xpath("./w:trPr/w:tblHeader"):
                    headers.append(r)
                for c, cell in enumerate(row.cells, 1+row.grid_cols_before):
                    key = cell._tc
                    if key in anchors:
                        anchor = anchors[key]
                        anchor.row_span = max(anchor.row_span, r-anchor.row+1)
                        anchor.column_span = max(anchor.column_span, c-anchor.column+1)
                        anchor.source.row_end = anchor.row+anchor.row_span-1
                        anchor.source.column_end = anchor.column+anchor.column_span-1
                    else:
                        anchor = Cell(row=r,column=c,text=cell.text,source=source.model_copy(update={"row_start":r,"row_end":r,"column_start":c,"column_end":c}))
                        anchors[key] = anchor
                        cells.append(anchor)
                        if cell.tables:
                            doc.warnings.append(f"Word 表 {table_index} 包含嵌套表：嵌套结构尚未进入索引，请查看原文")
            if block.rows and block.columns:
                _add_table(doc,len(block.rows),len(block.columns),cells,source,
                    header_rows=headers or [1],header_origin="explicit" if headers else "heuristic")
    if word.inline_shapes:
        doc.warnings.append("Word 内嵌图片未识别；OCR 接口目前仅定义 PDF 页/区域输入")
    doc.warnings.append("Word 来源按正文段落/表格定位；页眉页脚、文本框、批注和修订内容暂不解析")


def _scalar(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _excel(path, doc):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    try:
        for sheet in book:
            doc.sheets.append(sheet.title)
            # Ignore formatting-only empty cells; bound work using actual stored cells.
            occupied = [c for c in sheet._cells.values() if c.value is not None]
            if not occupied:
                continue
            min_r,max_r = min(c.row for c in occupied),max(c.row for c in occupied)
            min_c,max_c = min(c.column for c in occupied),max(c.column for c in occupied)
            for merge in sheet.merged_cells.ranges:
                min_r,max_r=min(min_r,merge.min_row),max(max_r,merge.max_row)
                min_c,max_c=min(min_c,merge.min_col),max(max_c,merge.max_col)
            if (max_r-min_r+1)*(max_c-min_c+1)>200000:
                raise ValueError(f"工作表 {sheet.title} 的有效区域超过 200000 个单元格，请拆分后上传")
            source = SourceLocation(format="xlsx",sheet=sheet.title,
                cell_range=f"{get_column_letter(min_c)}{min_r}:{get_column_letter(max_c)}{max_r}")
            merges={}
            covered=set()
            for m in sheet.merged_cells.ranges:
                merges[(m.min_row,m.min_col)]=m
                covered.update((r,c) for r in range(m.min_row,m.max_row+1) for c in range(m.min_col,m.max_col+1) if (r,c)!=(m.min_row,m.min_col))
            cells=[]
            for row in sheet.iter_rows(min_row=min_r,max_row=max_r,min_col=min_c,max_col=max_c):
                for cell in row:
                    if (cell.row,cell.column) in covered:
                        continue
                    if cell.data_type=="f" and not isinstance(cell.value,str):
                        value=getattr(cell.value,"text",None) or "[unsupported formula type]"
                        doc.warnings.append(f"{sheet.title}!{cell.coordinate} 为特殊公式；保留可用表达式，计算范围请核对原文件")
                    else:
                        value=_scalar(cell.value)
                    formula=value if cell.data_type=="f" else None
                    cv=_scalar(cached[sheet.title][cell.coordinate].value) if formula else None
                    text="" if value is None else str(value)
                    if formula:
                        text=f"{formula} [cached: {cv if cv is not None else 'missing'}; not recalculated]"
                    m=merges.get((cell.row,cell.column))
                    ref=str(m) if m else cell.coordinate
                    cells.append(Cell(row=cell.row-min_r+1,column=cell.column-min_c+1,
                        row_span=m.max_row-m.min_row+1 if m else 1,column_span=m.max_col-m.min_col+1 if m else 1,
                        text=text,value=None if formula else value,value_type=cell.data_type,formula=formula,cached_value=cv,
                        cache_status=("missing" if cv is None else "unverified") if formula else "not_formula",number_format=cell.number_format,
                        source=SourceLocation(format="xlsx",sheet=sheet.title,cell_range=ref)))
            # Leading full-width merged title/unit rows plus the next row are candidate headers.
            header_rows=[1]
            for r in range(1,min(5,max_r-min_r+1)):
                anchors=[c for c in cells if c.row==r and c.text]
                if len(anchors)==1 and anchors[0].column_span==max_c-min_c+1:
                    header_rows.append(r+1)
                else:
                    break
            _add_table(doc,max_r-min_r+1,max_c-min_c+1,cells,source,header_rows=header_rows,header_origin="heuristic")
            if sheet.sheet_state != "visible":
                doc.warnings.append(f"工作表 {sheet.title} 为隐藏状态，已按原始单元格解析")
        doc.warnings.append("Excel 保留原值、格式、合并区域与公式缓存；不执行公式，缓存新鲜度未知。每张工作表保存一个有效区域，开头合并标题及随后一行为候选表头；多块独立表与复杂表头需核验")
    finally:
        book.close()
        cached.close()
