"""Retrieval views only; complete cells and geometry stay in Document JSON."""
import hashlib
import json
import re
from .document import Document, location_label


def chunk_document(document: Document, tokenizer, max_tokens=300):
    if max_tokens < 16:
        raise ValueError("max_tokens must be >= 16")
    count = lambda text: len(tokenizer.encode(text, add_special_tokens=True, truncation=False, verbose=False))
    chunks = []
    tables = {t.table_id:t for t in document.tables}

    def emit(text, source, element_ids, **extra):
        if not text.strip():
            return
        raw = source.model_dump(exclude_none=True)
        digest=hashlib.sha256(json.dumps([text,raw,element_ids,extra],sort_keys=True,ensure_ascii=False).encode()).hexdigest()[:16]
        chunks.append(dict(chunk_id=f"{document.document_id}_{document.parser_version}_t{max_tokens}_c{len(chunks)+1}_{digest}",
            document_id=document.document_id,title=document.filename,page=source.page,text=text,
            word_count=len(text.split()),token_count=count(text),schema_version=document.schema_version,
            parser_version=document.parser_version,element_id=element_ids[0],context_refs=element_ids,
            source=raw,source_label=location_label(raw),
            ocr_pending_pages=sorted({r.page for r in document.ocr_requests}),**extra))

    def pieces(text, prefix=""):
        # Character boundaries preserve exact original text, including CJK and punctuation.
        offset=0
        while offset<len(text):
            lo,hi=1,len(text)-offset
            while lo<hi:
                mid=(lo+hi+1)//2
                if count(prefix+text[offset:offset+mid])<=max_tokens: lo=mid
                else: hi=mid-1
            if count(prefix+text[offset:offset+lo])>max_tokens:
                raise ValueError("Chunk context exceeds token budget")
            if offset+lo<len(text):
                # Prefer a complete whitespace-delimited item over cutting a financial number.
                boundary=max(text.rfind(" ",offset,offset+lo),text.rfind("\n",offset,offset+lo))
                if boundary>offset:
                    lo=boundary-offset+1
            # Token counts need not be monotonic when a word boundary changes.
            while lo and count(prefix+text[offset:offset+lo])>max_tokens:
                lo-=1
            if not lo:
                raise ValueError("Chunk context exceeds token budget")
            yield prefix+text[offset:offset+lo],offset,offset+lo
            offset+=lo

    pending=[]
    preceding=[]
    def flush():
        if not pending: return
        text=""; spans=[]
        for element in pending:
            start=len(text); text+=element.text+"\n"
            spans.append((start,len(text),element))
        for part,start,end in pieces(text):
            touched=[e for a,b,e in spans if a<end and b>start]
            source=touched[0].source.model_copy(deep=True)
            if source.format=="pdf":
                boxes=[e.source.bbox for e in touched]
                source.bbox=[min(b[0] for b in boxes),min(b[1] for b in boxes),max(b[2] for b in boxes),max(b[3] for b in boxes)]
            emit(part,source,[e.element_id for e in touched],
                text_spans=[dict(element_id=e.element_id,start=max(0,start-a),end=min(len(e.text),end-a))
                    for a,b,e in spans if a<end and b>start])
        pending.clear()

    elements=document.elements
    if document.format=="pdf":
        from .normalizer import normalize_pages
        pages=[dict(page=page,blocks=[dict(block_id=e.element_id,text=e.text,bbox=e.source.bbox)
            for e in elements if e.source.page==page and e.kind!="table"])
            for page in range(1,(document.page_count or 0)+1)]
        keep={b["block_id"] for page in normalize_pages(pages) for b in page["blocks"]}
        elements=[e for e in elements if e.kind=="table" or e.element_id in keep]
    for element in elements:
        if element.kind != "table":
            if pending and (element.source.page != pending[0].source.page or element.source.format != "pdf"):
                flush()
            pending.append(element)
            preceding.append(element)
            preceding[:]=preceding[-3:]
            continue
        flush()
        table=tables[element.table_id]
        headers=[c for c in table.cells if any(c.row<=r<c.row+c.row_span for r in table.header_rows)]
        header_text=" | ".join(f"C{i}: {name}" for i,name in enumerate(table.column_names,1)) if table.column_names else " | ".join(f"C{c.column}: {c.text}" for c in headers)
        # Explicit year labels travel with each value, including when a long row splits.
        years = {}
        for column in range(1, table.columns+1):
            labels = [c.text.strip() for c in headers if c.column <= column < c.column+c.column_span
                      and re.fullmatch(r'(?:19|20)\d{2}', c.text.strip())]
            if len(set(labels)) == 1:
                years[column] = labels[0]
        if table.column_names:
            for column,name in enumerate(table.column_names,1):
                if re.fullmatch(r'(?:19|20)\d{2}',name.strip()):
                    years[column] = name.strip()
        neighbors=[e for e in preceding if e.source.format==element.source.format
            and (e.source.format!="pdf" or e.source.page==element.source.page)]
        neighbor_text=" | ".join(e.text for e in neighbors)
        # Repeat nearby caption/unit text only when complete and within budget.
        nearby=f"Nearby text: {neighbor_text}\n" if neighbor_text and count(neighbor_text)<=max_tokens//5 else ""
        prefix=nearby+f"Table {table.table_id}; header ({table.header_origin}): {header_text}\n"
        # Huge headers remain separate indexed rows and in the structured table; no silent clipping.
        if count(prefix)>max_tokens//2:
            prefix=f"Table {table.table_id}; headers in structured source\n"
            nearby=""
        row_cells={}
        for cell in table.cells:
            for row in range(cell.row,cell.row+cell.row_span):
                row_cells.setdefault(row,[]).append(cell)
        for row in range(1,table.rows+1):
            cells=row_cells.get(row,[])
            if not any(c.text for c in cells): continue
            row_text=" | ".join(f"C{c.column}" + (f" [{years[c.column]}]" if c.column in years and c.column_span == 1 else "") + f": {c.text}" +
                (f" [format={c.number_format}]" if c.number_format and c.number_format!='General' else "") for c in cells)
            source=table.source.model_copy(deep=True)
            source.row_start=source.row_end=row
            if source.format=="xlsx":
                from openpyxl.utils.cell import range_boundaries, get_column_letter
                left,top,right,_=range_boundaries(table.source.cell_range)
                source.cell_range=f"{get_column_letter(left)}{top+row-1}:{get_column_letter(right)}{top+row-1}"
            for part,start,end in pieces(row_text,prefix):
                emit(part,source,[element.element_id]+([e.element_id for e in neighbors] if nearby else []),table_id=table.table_id,
                    cell_refs=[c.source.model_dump(exclude_none=True) for c in cells],
                    header_cell_refs=[c.source.model_dump(exclude_none=True) for c in headers]+[loc.model_dump(exclude_none=True) for loc in table.column_header_sources],
                    text_span=[start,end])
        preceding.clear()
    flush()
    return chunks
