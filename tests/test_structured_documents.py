"""Offline structural fixtures, no OCR/embedding/LLM calls."""
import json
from pathlib import Path
import pymupdf
import pytest
from docx import Document as Word
from openpyxl import Workbook
from financial_rag.ingestion.adapters import parse_document
from financial_rag.ingestion.document import SourceLocation
from financial_rag.ingestion.structured_chunker import chunk_document
from financial_rag.memory.evidence import EvidenceMemory


class Tokenizer:
    def encode(self, text, add_special_tokens=True, **kwargs):
        return list(text) + ([0,0] if add_special_tokens else [])


def word_file(path):
    doc=Word(); doc.add_heading('Results',level=1);doc.add_paragraph('Revenue and units: USD millions')
    table=doc.add_table(rows=3,cols=3)
    table.cell(0,0).text='Metric';table.cell(0,1).text='2024';table.cell(0,2).text='2025'
    table.cell(1,0).merge(table.cell(2,0)).text='Revenue'
    table.cell(1,1).merge(table.cell(1,2)).text='100'
    table.cell(2,1).text='200';table.cell(2,2).text='300'
    doc.add_paragraph('After table');doc.save(path)
    return path


def excel_file(path):
    book=Workbook();sheet=book.active;sheet.title='Income statement'
    sheet['B2']='USD millions';sheet.merge_cells('B2:D2')
    sheet['B3']='Metric';sheet['C3']=2024;sheet['D3']=2025
    sheet['B4']='Margin';sheet['C4']=0;sheet['C4'].number_format='0.0%';sheet['D4']='-'
    sheet['B5']='Total';sheet['C5']='=SUM(C4:C4)'
    book.create_sheet('Notes')['A1']='Source note'
    book.save(path);return path


def test_word_order_merges_source_and_chunks(tmp_path):
    doc=parse_document(word_file(tmp_path/'report.docx'))
    assert [e.kind for e in doc.elements]==['heading','paragraph','table','paragraph']
    t=doc.tables[0]
    assert next(c for c in t.cells if c.text=='Revenue').row_span==2
    assert next(c for c in t.cells if c.text=='100').column_span==2
    assert len(t.cells)==7
    assert t.source.heading_path==['Results'] and t.source.page is None
    chunks=chunk_document(doc,Tokenizer(),100)
    assert all(c['token_count']<=100 and c['page'] is None for c in chunks)
    row=next(c for c in chunks if c.get('table_id') and 'Revenue' in c['text'])
    assert row['table_id']=='t1' and row['header_cell_refs']
    memory=EvidenceMemory();memory.add_evidence([row],{row['chunk_id']:row})
    assert memory.citation_map['E1']['source']['format']=='docx'
    assert 'Results' in memory.build_context()
    bad=dict(row,source={**row['source'],'table':2})
    with pytest.raises(ValueError):memory.add_evidence([bad],{bad['chunk_id']:bad})


def test_excel_formula_blank_zero_dash_and_merge(tmp_path):
    doc=parse_document(excel_file(tmp_path/'report.xlsx'))
    assert doc.sheets==['Income statement','Notes']
    table=doc.tables[0]
    assert table.source.cell_range=='B2:D5'
    assert table.header_rows==[1,2]
    cells={c.source.cell_range:c for c in table.cells}
    assert cells['B2:D2'].column_span==3
    assert cells['C4'].value==0 and cells['C4'].number_format=='0.0%'
    assert cells['D4'].value=='-' and cells['D5'].value is None
    assert cells['C5'].formula=='=SUM(C4:C4)' and cells['C5'].cached_value is None
    assert cells['C5'].cache_status=='missing'
    chunks=chunk_document(doc,Tokenizer(),200)
    assert any(c['source']['cell_range']=='B4:D4' and 'Margin' in c['text'] for c in chunks)
    assert all(c['page'] is None and c['token_count']<=200 for c in chunks)
    assert json.loads(doc.model_dump_json())['tables'][0]['cells']


def test_pdf_table_geometry_no_duplicate_body_and_ocr(tmp_path):
    path=tmp_path/'table.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        page.insert_text((50,40),'Revenue USD millions')
        for y in [60,90,120,150]:page.draw_line((50.35965,y),(350.35965,y))
        for x in [50.35965,150.35965,250.35965,350.35965]:page.draw_line((x,60),(x,150))
        for row,values in enumerate([['Metric','2024','2025'],['Revenue','100','200'],['Profit','10','20']]):
            for col,value in enumerate(values):page.insert_text((55+col*100,80+row*30),value)
        pdf.new_page()  # empty page must not silently be treated as fully read
        pdf.save(path)
    doc=parse_document(path)
    assert len(doc.tables)==1
    assert doc.tables[0].rows==3 and len(doc.tables[0].cells)==9
    assert all(c.source.page==1 and c.source.bbox for c in doc.tables[0].cells)
    assert not any(e.text=='100' for e in doc.elements)
    assert doc.ocr_requests[0].page==2 and doc.ocr_requests[0].status=='needs_ocr'
    chunks=chunk_document(doc,Tokenizer(),200)
    assert any(c.get('table_id') and '100' in c['text'] for c in chunks)
    assert all(c['source']['bbox'] and c['page']==1 for c in chunks)
    assert parse_document(path).document_id==doc.document_id


def test_long_cells_preserve_content_without_embedding_truncation(tmp_path):
    path=tmp_path/'long.docx';doc=Word();t=doc.add_table(rows=2,cols=1)
    t.cell(0,0).text='Header';t.cell(1,0).text='中文ABC'*200;doc.save(path)
    parsed=parse_document(path);chunks=chunk_document(parsed,Tokenizer(),100)
    row=[c for c in chunks if c.get('source',{}).get('row_start')==2]
    assert len(row)>1 and all(c['token_count']<=100 for c in chunks)
    assert ''.join(c['text'].split('\n',1)[1] for c in row)=='C1: '+'中文ABC'*200
    assert parsed.tables[0].cells[1].text=='中文ABC'*200


@pytest.mark.parametrize('source',[dict(format='docx',page=1,paragraph=1),dict(format='pdf'),dict(format='xlsx',sheet='S')])
def test_invalid_locations_rejected(source):
    with pytest.raises(ValueError):SourceLocation(**source)


def test_provenance_survives_dense_sparse_fusion_citation_and_sqlite(tmp_path):
    import numpy as np
    from financial_rag.indexing.vector_store import VectorStore
    from financial_rag.indexing.bm25_index import BM25Index
    from financial_rag.retrieval.dense import dense_search
    from financial_rag.retrieval.sparse import bm25_search
    from financial_rag.retrieval.fusion import reciprocal_rank_fusion
    from financial_rag.generation.citations import check_citations
    chunks=chunk_document(parse_document(excel_file(tmp_path/'report.xlsx')),Tokenizer(),300)
    vectors=np.tile([1.,0.],(len(chunks),1));store=VectorStore(tmp_path/'vectors.sqlite3')
    store.replace_all(chunks,vectors,'fake');restored,matrix=store.load('fake')
    assert restored==chunks
    dense=dense_search(np.array([1.,0.]),matrix,restored,20)
    sparse=bm25_search('Margin',BM25Index(restored),20)
    fused=reciprocal_rank_fusion(dense,sparse,top_k=20)
    assert all(i['source']['format']=='xlsx' for i in dense+sparse+fused)
    memory=EvidenceMemory();memory.add_evidence(fused,{c['chunk_id']:c for c in chunks})
    citation=check_citations('value [E1]',memory.citation_map)['sources'][0]
    assert citation['page'] is None and citation['cell_refs'] and citation['source_label']


def test_excel_cached_formula_preserved_without_recalculation(tmp_path):
    from zipfile import ZipFile,ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    path=excel_file(tmp_path/'cache.xlsx')
    with ZipFile(path) as z:parts={name:z.read(name) for name in z.namelist()}
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    xml.find(".//s:c[@r='C5']/s:v",ns).text='17'
    parts['xl/worksheets/sheet1.xml']=ET.tostring(xml)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    cell=next(c for c in parse_document(path).tables[0].cells if c.formula)
    assert cell.cached_value==17 and cell.cache_status=='unverified'
    assert cell.formula=='=SUM(C4:C4)'


def test_source_snapshots_are_deep_copies(tmp_path):
    chunks=chunk_document(parse_document(word_file(tmp_path/'copy.docx')),Tokenizer(),300)
    row=next(c for c in chunks if c.get('table_id'))
    memory=EvidenceMemory();memory.add_evidence([row],{row['chunk_id']:row})
    memory.citation_map['E1']['source']['table']=999
    assert memory.citation_map['E1']['source']['table']==1


def test_pdf_footer_kept_in_original_not_retrieval(tmp_path):
    path=tmp_path/'footer.pdf'
    with pymupdf.open() as pdf:
        for i in range(3):
            page=pdf.new_page();page.insert_text((50,50),'Revenue 123')
            page.insert_text((50,800),f'Company | 2024 Form 10-K | {i+1}')
        pdf.save(path)
    doc=parse_document(path)
    assert any('Form 10-K' in e.text for e in doc.elements)
    assert all('Form 10-K' not in c['text'] for c in chunk_document(doc,Tokenizer()))


def test_pdf_uneven_right_edges_do_not_create_phantom_columns(tmp_path):
    path=tmp_path/'uneven.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        page.draw_rect((50.35965,60,350.35965,90))
        for r in range(2):
            for c in range(2):
                page.draw_rect((50.35965+c*145,90+r*30,50.35965+(c+1)*145,120+r*30))
                page.insert_text((55+c*145,110+r*30),str(10+r+c))
        page.insert_text((55,80),'Merged header')
        pdf.save(path)
    parsed=parse_document(path)
    assert parsed.tables
    table=parsed.tables[0]
    assert table.columns==2
    assert next(c for c in table.cells if 'Merged header' in c.text).column_span==2



def test_borderless_year_alignment_with_baseline_jitter(tmp_path):
    path=tmp_path/'years.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        page.insert_text((40,40),'Balance sheet (USD millions)')
        page.insert_text((40,60),'At December 31,')
        page.insert_text((380,60),'2022')
        page.insert_text((480,60),'2021')
        for i,(label,a,b) in enumerate([('Inventories','2,388','3,055'),
                ('Prepaid expenses','8,358','6,906'),('Total current assets','37,857','36,728'),
                ('Total current liabilities','50,171','47,160')]):
            y=85+i*25
            page.insert_text((40,y+1.7),label)
            # Right aligned numeric columns; previous-year baseline comes first in y sort.
            for x,text,dy in [(404.5,a,.009),(504.5,b,0)]:
                page.insert_text((x-pymupdf.get_text_length(text,fontsize=11),y+dy),text)
        pdf.save(path)
    doc=parse_document(path)
    assert len(doc.tables)==1
    table=doc.tables[0]
    last={c.column:c for c in table.cells if c.row==5}
    assert last[2].text=='50,171' and last[3].text=='47,160'
    chunks=chunk_document(doc,Tokenizer(),300)
    target=next(c for c in chunks if 'Total current liabilities' in c['text'])
    assert 'C2 [2022]: 50,171' in target['text']
    assert 'C3 [2021]: 47,160' in target['text']
    assert 'USD millions' in target['text']
    assert target['header_cell_refs'] and target['cell_refs']
    assert doc.parser_version=='native-v2'


def test_borderless_prose_and_unaligned_numbers_not_guessed(tmp_path):
    path=tmp_path/'prose.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        page.insert_text((40,50),'Revenue in 2022 increased over 2021.')
        for i in range(5):page.insert_text((40,80+i*20),f'Note {i}: revenue was 100 and 200.')
        pdf.save(path)
    doc=parse_document(path)
    assert not doc.tables
    assert any('Revenue in 2022' in e.text for e in doc.elements)


def test_excel_year_value_unit_associations_survive_chunks(tmp_path):
    doc=parse_document(excel_file(tmp_path/'years.xlsx'))
    chunks=chunk_document(doc,Tokenizer(),300)
    target=next(c for c in chunks if 'Margin' in c['text'])
    assert 'USD millions' in target['text']
    assert 'C2 [2024]: 0 [format=0.0%]' in target['text']
    assert 'C3 [2025]: -' in target['text']
    assert target['source']['cell_range']=='B4:D4'


def test_word_years_preserve_negative_decimal_values(tmp_path):
    path=tmp_path/'years.docx';doc=Word();doc.add_paragraph('USD millions')
    table=doc.add_table(rows=2,cols=3)
    for i,v in enumerate(['Metric','2022','2021']):table.cell(0,i).text=v
    for i,v in enumerate(['Profit','(12.50)','0']):table.cell(1,i).text=v
    doc.save(path)
    chunks=chunk_document(parse_document(path),Tokenizer(),300)
    row=next(c for c in chunks if 'Profit' in c['text'])
    assert 'C2 [2022]: (12.50)' in row['text'] and 'C3 [2021]: 0' in row['text']
    assert 'USD millions' in row['text']



def test_borderless_currency_symbols_are_not_dropped(tmp_path):
    path=tmp_path/'currency.pdf'
    with pymupdf.open() as pdf:
        page=pdf.new_page()
        page.insert_text((380,60),'2022');page.insert_text((480,60),'2021')
        for i in range(3):
            y=90+i*25
            page.insert_text((40,y),f'Metric {i}')
            for x in (380,480):
                page.insert_text((x-12,y),'$')
                page.insert_text((x,y),'100')
        pdf.save(path)
    doc=parse_document(path)
    assert len(doc.tables)==1
    assert len([c for c in doc.tables[0].cells if c.text=='$ 100'])==6
